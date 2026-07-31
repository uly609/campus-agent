from __future__ import annotations

import base64
import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.llm.router import ProviderRouter
from app.multimodal.image_attributes import analyze_chat_image

MAX_FILE_BYTES = 10 * 1024 * 1024
ROW_CHUNK_SIZE = 60
MAX_TEXT_CHUNK_CHARS = 3000

ALLOWED_EXTENSIONS = frozenset(
    {".xlsx", ".csv", ".pdf", ".txt", ".md", ".markdown", ".png", ".jpg", ".jpeg", ".webp"}
)
IMAGE_MIME_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


@dataclass(frozen=True)
class ParsedFileChunk:
    chunk_id: str
    kind: str
    title: str
    text: str
    metadata: dict[str, str]


@dataclass(frozen=True)
class FileParseResult:
    file_name: str
    kind: str
    chunks: list[ParsedFileChunk]
    warnings: list[str]


def _validate_input(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError(f"不支持的文件类型 {suffix or '(无扩展名)'}")
    if not content:
        raise ValueError("文件内容为空")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("文件超过 10MB 限制")
    return suffix


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _markdown_table(header: list[str], rows: list[list[str]]) -> str:
    columns = max(len(header), max((len(row) for row in rows), default=0))
    normalized = header + [""] * (columns - len(header))
    lines = ["| " + " | ".join(normalized) + " |", "| " + " | ".join("---" for _ in range(columns)) + " |"]
    for row in rows:
        cells = [_clean_cell(cell) for cell in row] + [""] * (columns - len(row))
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def _parse_xlsx(filename: str, content: bytes) -> list[ParsedFileChunk]:
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover - dependency is declared in requirements
        raise ValueError("缺少 openpyxl 依赖，无法解析 Excel") from exc

    chunks: list[ParsedFileChunk] = []
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [_clean_cell(cell) for cell in row]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        header = rows[0]
        for start in range(1, len(rows), ROW_CHUNK_SIZE):
            block = rows[start : start + ROW_CHUNK_SIZE]
            end = min(start + ROW_CHUNK_SIZE, len(rows)) - 1
            chunks.append(
                ParsedFileChunk(
                    chunk_id=f"{filename}:{sheet.title}:{start}:{end}",
                    kind="table",
                    title=f"{filename} / {sheet.title}",
                    text=_markdown_table(header, block),
                    metadata={
                        "file_name": filename,
                        "sheet_name": sheet.title,
                        "row_start": str(start + 1),
                        "row_end": str(end + 1),
                        "kind": "table",
                    },
                )
            )
    workbook.close()
    return chunks


def _parse_csv(filename: str, content: bytes) -> list[ParsedFileChunk]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [[_clean_cell(cell) for cell in row] for row in reader if any(_clean_cell(cell) for cell in row)]
    if not rows:
        return []
    header = rows[0]
    chunks: list[ParsedFileChunk] = []
    for start in range(1, len(rows), ROW_CHUNK_SIZE):
        block = rows[start : start + ROW_CHUNK_SIZE]
        end = min(start + ROW_CHUNK_SIZE, len(rows)) - 1
        chunks.append(
            ParsedFileChunk(
                chunk_id=f"{filename}:csv:{start}:{end}",
                kind="table",
                title=f"{filename} / CSV",
                text=_markdown_table(header, block),
                metadata={
                    "file_name": filename,
                    "sheet_name": "CSV",
                    "row_start": str(start + 1),
                    "row_end": str(end + 1),
                    "kind": "table",
                },
            )
        )
    return chunks


def _parse_pdf(filename: str, content: bytes) -> tuple[list[ParsedFileChunk], list[str]]:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency is declared in requirements
        raise ValueError("缺少 pdfplumber 依赖，无法解析 PDF") from exc

    chunks: list[ParsedFileChunk] = []
    warnings: list[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            tables: list[str] = []
            for table in page.extract_tables() or []:
                cleaned = [
                    [_clean_cell(cell) for cell in row]
                    for row in table
                    if any(_clean_cell(cell) for cell in row)
                ]
                if cleaned:
                    tables.append(_markdown_table(cleaned[0], cleaned[1:]))
            page_text = "\n\n".join(part for part in (text.strip(), "\n\n".join(tables)) if part.strip())
            if not page_text.strip():
                continue
            chunks.append(
                ParsedFileChunk(
                    chunk_id=f"{filename}:page:{index}",
                    kind="table" if tables and len(text.strip()) < 200 else "text",
                    title=f"{filename} / 第 {index} 页",
                    text=page_text[:MAX_TEXT_CHUNK_CHARS * 2],
                    metadata={
                        "file_name": filename,
                        "page": str(index),
                        "pages": str(len(pdf.pages)),
                        "kind": "table" if tables else "text",
                    },
                )
            )
        if not chunks:
            warnings.append("PDF 未提取到文本，可能需要扫描件 OCR 或重新导出")
    return chunks, warnings


def _parse_text(filename: str, content: bytes) -> list[ParsedFileChunk]:
    text = content.decode("utf-8", errors="replace").strip()
    if not text:
        return []
    paragraphs = [paragraph.strip() for paragraph in re.split(r"\n\s*\n", text) if paragraph.strip()]
    chunks: list[ParsedFileChunk] = []
    buffer: list[str] = []
    buffer_size = 0
    for paragraph in paragraphs:
        if buffer_size + len(paragraph) > MAX_TEXT_CHUNK_CHARS and buffer:
            chunks.append(
                ParsedFileChunk(
                    chunk_id=f"{filename}:text:{len(chunks)}",
                    kind="text",
                    title=filename,
                    text="\n\n".join(buffer)[:MAX_TEXT_CHUNK_CHARS],
                    metadata={"file_name": filename, "kind": "text"},
                )
            )
            buffer = []
            buffer_size = 0
        buffer.append(paragraph)
        buffer_size += len(paragraph)
    if buffer:
        chunks.append(
            ParsedFileChunk(
                chunk_id=f"{filename}:text:{len(chunks)}",
                kind="text",
                title=filename,
                text="\n\n".join(buffer)[:MAX_TEXT_CHUNK_CHARS],
                metadata={"file_name": filename, "kind": "text"},
            )
        )
    return chunks


async def _parse_image(
    filename: str,
    content: bytes,
    suffix: str,
    router: ProviderRouter | None,
) -> tuple[ParsedFileChunk, list[str]]:
    data_url = f"data:{IMAGE_MIME_TYPES[suffix]};base64,{base64.b64encode(content).decode('ascii')}"
    attributes = await analyze_chat_image(data_url, router)
    warnings: list[str] = []
    if not attributes:
        warnings.append("视觉模型未返回结构化图片信息")
    analysis = attributes.get("_analysis", {})
    degraded = bool(analysis.get("degraded"))
    if degraded:
        warnings.append("图片识别使用离线演示模型，不代表真实视觉理解")
    text = "图片解析结果：\n"
    text += f"- 摘要：{attributes.get('summary', '未提取到摘要')}\n"
    for key, label in (
        ("category", "类别"),
        ("color", "颜色"),
        ("brand", "品牌"),
        ("material", "材质"),
        ("location_hints", "地点线索"),
        ("visible_text", "可见文字"),
        ("safety_flags", "安全标记"),
    ):
        value = attributes.get(key)
        if value:
            text += f"- {label}：{value if isinstance(value, str) else '、'.join(map(str, value))}\n"
    return (
        ParsedFileChunk(
            chunk_id=f"{filename}:image",
            kind="image",
            title=f"{filename} / 图片识别",
            text=text.strip(),
            metadata={
                "file_name": filename,
                "kind": "image",
                "provider": str(analysis.get("provider", "")),
                "model": str(analysis.get("model", "")),
                "degraded": "true" if degraded else "false",
            },
        ),
        warnings,
    )


async def parse_document_file(
    filename: str,
    content: bytes,
    router: ProviderRouter | None = None,
) -> FileParseResult:
    suffix = _validate_input(filename, content)
    warnings: list[str] = []
    if suffix == ".xlsx":
        chunks = _parse_xlsx(filename, content)
        kind = "table"
    elif suffix == ".csv":
        chunks = _parse_csv(filename, content)
        kind = "table"
    elif suffix == ".pdf":
        chunks, warnings = _parse_pdf(filename, content)
        kind = "pdf"
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        chunk, warnings = await _parse_image(filename, content, suffix, router)
        chunks = [chunk]
        kind = "image"
    else:
        chunks = _parse_text(filename, content)
        kind = "text"
    if not chunks:
        raise ValueError("文件解析后没有可用的文本片段")
    return FileParseResult(file_name=filename, kind=kind, chunks=chunks, warnings=warnings)
